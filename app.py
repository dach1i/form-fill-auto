import streamlit as st
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from datetime import datetime
import io
import os
import copy
import json
import subprocess
from subprocess import PIPE, run
import tempfile
import zipfile
import shutil
from pathlib import Path
import re
from PIL import Image

def convert_doc_to_pdf_native(doc_file: Path, output_dir: Path = Path("."), timeout: int = 60):
    """Converts a doc file to pdf using libreoffice directly in headless mode.
    Implementation from Franky1/Streamlit-docx-converter.
    """
    exception = None
    output = None
    try:
        process = run(
            ['soffice', '--headless', '--convert-to', 'pdf:writer_pdf_Export', '--outdir', str(output_dir.resolve()), str(doc_file.resolve())],
            stdout=PIPE, stderr=PIPE,
            timeout=timeout, check=True
        )
        stdout = process.stdout.decode("utf-8")
        re_filename = re.search(r'-> (.*?) using filter', stdout)
        if re_filename:
            output = Path(re_filename[1]).resolve()
        else:
            cand = output_dir.joinpath(doc_file.stem + ".pdf")
            if cand.exists():
                output = cand.resolve()
    except Exception as e:
        exception = e
    return (output, exception)

st.set_page_config(page_title="DPR Generator", page_icon="🏗️", layout="wide")
st.title("🏗️ ყოველდღიური რეპორტის გენერატორი (DPR)")
st.caption("სითი მოლი საბურთალო — სრული ავტომატიზაცია")

target_file = None
for fname in ["template.docx", "template_2.docx", "CMC-CMS-DPR-007-20261008-Rev00.docx"]:
    if os.path.exists(fname):
        target_file = fname
        break

if not target_file:
    st.error("ვერ მოიძებნა 'template.docx'!")
    st.stop()

# ==================== AUTOSAVE DRAFT MANAGER ====================
AUTOSAVE_FILE = "dpr_autosave.json"

def load_draft():
    if os.path.exists(AUTOSAVE_FILE):
        try:
            with open(AUTOSAVE_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return {}
    return {}

def save_draft(data):
    try:
        with open(AUTOSAVE_FILE, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
    except Exception:
        pass

draft = load_draft()

with st.sidebar:
    st.header("⚙️ მართვის პანელი")
    
    col_sb1, col_sb2 = st.columns(2)
    if col_sb1.button("⚡ Auto-Fill Sample", use_container_width=True):
        sample_data = {
            "rep_num": "008",
            "rep_date": datetime.today().strftime("%Y-%m-%d"),
            "prepared_by": "მურად ახალაძე",
            "custom_doc_code": f"CMC-CMS-DPR-008-{datetime.today().strftime('%Y%m%d')}-Rev00",
            "client_name": "შპს „აი ჯი დეველოფმენტ ჯორჯია“",
            "site_location": "ქ.თბილისი ქავთარაძის ქ. №1",
            "t9": "+21°C", "w9": "1.8 მ/წ", "cond9": "მზიანი/sunny",
            "t14": "+24°C", "w14": "2.1 მ/წ", "cond14": "ნაწ. მოღრ./p. cloudy",
            "t18": "+22°C", "w18": "1.5 მ/წ", "cond18": "მზიანი/sunny",
            "task_count": 2,
            "tasks": [
                {
                    "num": 1,
                    "contractor": "ბკ ქონსთრაქშენი",
                    "manpower": "12",
                    "wbs": "WBS-01",
                    "name": "შპუნტების მონტაჟი",
                    "desc": "მიმდინარეობს E ბლოკში შპუნტების მონტაჟი"
                },
                {
                    "num": 2,
                    "contractor": "ბკ ქონსთრაქშენი",
                    "manpower": "8",
                    "wbs": "WBS-02",
                    "name": "ხიმინჯის გაბურღვა",
                    "desc": "მიმდინარეობს E ბლოკში ხიმინჯების გაბურღვა"
                }
            ],
            "photo_slots_count": 6,
            "captions": {
                "0": "სურ./Pic. 1 შპუნტების მონტაჟი",
                "1": "სურ./Pic. 2 შპუნტების მონტაჟი"
            },
            "staff_count": 3,
            "staff": [
                {
                    "num": 1,
                    "comp": "სიემსი",
                    "pos": "პროექტის მენეჯერი",
                    "name": "ავთო ჯაბაური",
                    "phone": "577 252 282",
                    "mail": "a.jabauri@cmc.ge"
                },
                {
                    "num": 2,
                    "comp": "სიემსი",
                    "pos": "ობიექტის ზედამხედველი",
                    "name": "მურად ახალაძე",
                    "phone": "598 109 291",
                    "mail": "m.akhaladze@cmc.ge"
                },
                {
                    "num": 3,
                    "comp": "ბკ ქონსთრაქშენი",
                    "pos": "ობიექტის ზედამხედველი",
                    "name": "ლაშა სამხარაძე",
                    "phone": "598 250 660",
                    "mail": "l.samkharadze@bkconstruction.ge"
                }
            ]
        }
        save_draft(sample_data)
        st.session_state.clear()
        st.rerun()

    if col_sb2.button("🔄 Reset Draft", use_container_width=True):
        if os.path.exists(AUTOSAVE_FILE):
            os.remove(AUTOSAVE_FILE)
        st.session_state.clear()
        st.rerun()

    st.caption("ყველა ცვლილება ინახება ავტომატურად.")

if "docx_bytes" not in st.session_state:
    st.session_state["docx_bytes"] = None
    st.session_state["pdf_bytes"] = None
    st.session_state["doc_code"] = None

def clone_row(table, source_row):
    tr_copy = copy.deepcopy(source_row._tr)
    for t in tr_copy.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'):
        t.text = ''
    table._tbl.append(tr_copy)
    return table.rows[-1]

def clear_cell_completely(cell):
    tc = cell._tc
    for child in list(tc):
        if not child.tag.endswith('tcPr'):
            tc.remove(child)
    cell.add_paragraph()

def safe_replace(p, old, new):
    if old in p.text:
        for r in p.runs:
            if old in r.text:
                r.text = r.text.replace(old, str(new))
        if old in p.text:
            p.text = p.text.replace(old, str(new))

def populate_cell(cell, text, font_size=8.0, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER):
    clear_cell_completely(cell)
    try:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    except Exception:
        pass
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(str(text) if text is not None else "")
    r.font.size = Pt(font_size)
    r.font.bold = bold
    r.font.color.rgb = RGBColor(0, 0, 0)

WEATHER_ICON_MAP = {
    "მზიანი/sunny": "word/media/image3.png",
    "ნაწ. მოღრ./p. cloudy": "word/media/image4.png",
    "მოღრუბლ./cloudy": "word/media/image5.png",
    "წვიმიანი/rainy": "word/media/image6.png",
    "თოვლიანი/snowy": "word/media/image7.png",
}

def load_weather_icons(docx_path):
    icons = {}
    try:
        with zipfile.ZipFile(docx_path, 'r') as z:
            for cond, arc in WEATHER_ICON_MAP.items():
                if arc in z.namelist():
                    icons[cond] = z.read(arc)
    except Exception:
        pass
    return icons

def set_weather_icon_cell(cell, condition_str, icons_dict):
    clear_cell_completely(cell)
    try:
        tcPr = cell._tc.get_or_add_tcPr()
        for child in list(tcPr):
            if child.tag.endswith('shd'):
                tcPr.remove(child)
        from docx.oxml import parse_xml
        from docx.oxml.ns import nsdecls
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="262626"/>')
        tcPr.append(shd)
    except Exception:
        pass
    try:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    except Exception:
        pass
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0

    img_data = icons_dict.get(condition_str) if icons_dict else None
    if img_data:
        r = p.add_run()
        r.add_picture(io.BytesIO(img_data), width=Inches(1.15))

def insert_centered_picture(cell, file_bytes, width=Inches(3.15)):
    clear_cell_completely(cell)
    try:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    except Exception:
        pass
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)

    max_w_in = 3.4
    max_h_in = 2.0
    try:
        with Image.open(io.BytesIO(file_bytes)) as img:
            w_px, h_px = img.size
            aspect = h_px / w_px
            if aspect > (max_h_in / max_w_in):
                target_height = Inches(max_h_in)
                target_width = Inches(max_h_in / aspect)
            else:
                target_width = Inches(max_w_in)
                target_height = Inches(max_w_in * aspect)
    except Exception:
        target_width = Inches(max_w_in)
        target_height = None

    r = p.add_run()
    if target_height:
        r.add_picture(io.BytesIO(file_bytes), width=target_width, height=target_height)
    else:
        r.add_picture(io.BytesIO(file_bytes), width=target_width)

def set_clean_wind_cell(cell, wind_val):
    clear_cell_completely(cell)
    try:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    except Exception:
        pass
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.0
    lbl = p.add_run("ქარის სიჩქარე Wind Speed:\n")
    lbl.font.size = Pt(7.5)
    lbl.font.bold = False
    lbl.font.color.rgb = RGBColor(0, 0, 0)
    val = p.add_run(str(wind_val if wind_val else "-"))
    val.font.size = Pt(8.5)
    val.font.bold = True
    val.font.color.rgb = RGBColor(0, 0, 0)

WEATHER_CONDITIONS = [
    "",
    "მზიანი/sunny",
    "ნაწ. მოღრ./p. cloudy",
    "მოღრუბლ./cloudy",
    "წვიმიანი/rainy",
    "თოვლიანი/snowy"
]

# ==================== INPUT FIELDS ====================
# 1. ძირითადი ინფორმაცია
st.subheader("1. ძირითადი ინფორმაცია / General Info")
col1, col2, col3, col4 = st.columns([1, 1, 1.5, 2])

rep_num = col1.text_input("რეპორტის №", value=draft.get("rep_num", ""))

raw_saved_date = draft.get("rep_date")
default_date = datetime.today().date()
if raw_saved_date:
    try:
        default_date = datetime.strptime(raw_saved_date, "%Y-%m-%d").date()
    except Exception:
        pass
rep_date = col2.date_input("თარიღი", value=default_date)
prepared_by = col3.text_input("მოამზადა / Prepared by", value=draft.get("prepared_by", ""))
custom_doc_code = col4.text_input("დოკუმენტის კოდი / Doc Code", value=draft.get("custom_doc_code", ""))

col_c1, col_c2 = st.columns(2)
client_name = col_c1.text_input("დამკვეთი / Client", value=draft.get("client_name", ""))
site_location = col_c2.text_input("ადგილმდებარეობა / Location", value=draft.get("site_location", ""))

st.divider()

# 2. ამინდი
st.subheader("2. ამინდი / Weather")
w1, w2, w3 = st.columns(3)

w1.markdown("#### 09:00")
t9 = w1.text_input("ტემპერატურა (09:00)", value=draft.get("t9", ""))
w9 = w1.text_input("ქარი (09:00)", value=draft.get("w9", ""))
saved_cond9 = draft.get("cond9", "")
cond9 = w1.selectbox("ამინდის პირობა (09:00)", WEATHER_CONDITIONS, index=WEATHER_CONDITIONS.index(saved_cond9) if saved_cond9 in WEATHER_CONDITIONS else 0)

w2.markdown("#### 14:00")
t14 = w2.text_input("ტემპერატურა (14:00)", value=draft.get("t14", ""))
w14 = w2.text_input("ქარი (14:00)", value=draft.get("w14", ""))
saved_cond14 = draft.get("cond14", "")
cond14 = w2.selectbox("ამინდის პირობა (14:00)", WEATHER_CONDITIONS, index=WEATHER_CONDITIONS.index(saved_cond14) if saved_cond14 in WEATHER_CONDITIONS else 0)

w3.markdown("#### 18:00")
t18 = w3.text_input("ტემპერატურა (18:00)", value=draft.get("t18", ""))
w18 = w3.text_input("ქარი (18:00)", value=draft.get("w18", ""))
saved_cond18 = draft.get("cond18", "")
cond18 = w3.selectbox("ამინდის პირობა (18:00)", WEATHER_CONDITIONS, index=WEATHER_CONDITIONS.index(saved_cond18) if saved_cond18 in WEATHER_CONDITIONS else 0)

st.divider()

# 3. სამუშაოები (Daily Works)
st.subheader("3. სამუშაოების აღწერა / Daily Works")
saved_task_count = int(draft.get("task_count", 1))
task_count = st.number_input("სამუშაოების რაოდენობა", min_value=1, max_value=25, value=saved_task_count)

draft_tasks = draft.get("tasks", [])
tasks = []
for i in range(int(task_count)):
    st.markdown(f"**სამუშაო #{i+1}**")
    tc1, tc2, tc3, tc4, tc5 = st.columns([2, 1, 1, 2, 4])
    
    t_prev = draft_tasks[i] if i < len(draft_tasks) else {}
    w_name = tc1.text_input(f"დასახელება #{i+1}", value=t_prev.get("name", ""), key=f"wn_{i}")
    m_pow = tc2.text_input(f"პერსონალი #{i+1}", value=t_prev.get("manpower", ""), key=f"mp_{i}")
    wbs_c = tc3.text_input(f"WBS #{i+1}", value=t_prev.get("wbs", ""), key=f"wbs_{i}")
    c_tor = tc4.text_input(f"კონტრაქტორი #{i+1}", value=t_prev.get("contractor", ""), key=f"co_{i}")
    desc = tc5.text_input(f"აღწერა #{i+1}", value=t_prev.get("desc", ""), key=f"de_{i}")
    tasks.append({"num": i+1, "contractor": c_tor, "manpower": m_pow, "wbs": wbs_c, "name": w_name, "desc": desc})

st.divider()

# 4. ფოტოები
st.subheader("4. სამუშაოების ფოტომასალა / Photos & Descriptions")
st.info("💡 თითო გვერდზე თავსდება 6 ფოტო. 7+ ფოტოს შემთხვევაში ავტომატურად შეიქმნება ახალი გვერდი (მაქსიმუმ 100 ფოტო).")

saved_photo_slots = min(int(draft.get("photo_slots_count", 6)), 100)
photo_slots_count = st.number_input("რამდენი ფოტოს ატვირთვა გსურთ?", min_value=1, max_value=100, value=saved_photo_slots)
draft_caps = draft.get("captions", {})
photos_data = []

for pi in range(int(photo_slots_count)):
    p_c1, p_c2 = st.columns([2, 3])
    up_file = p_c1.file_uploader(f"ფოტო #{pi+1}", type=["jpg", "jpeg", "png"], key=f"photo_up_{pi}")
    def_cap = draft_caps.get(str(pi), "")
    p_caption = p_c2.text_input(f"აღწერა #{pi+1}", value=def_cap, key=f"photo_cap_{pi}")
    draft_caps[str(pi)] = p_caption
    if up_file is not None:
        photos_data.append({"file": up_file, "caption": p_caption, "num": pi + 1})

st.divider()

# 5. პერსონალი (Key Staff)
st.subheader("5. სამშენებლო მოედანზე მომუშავე პერსონალი / Key Staff")
st.caption("ჩაწერეთ მხოლოდ დღეს მყოფი თანამშრომლები.")
saved_staff_count = int(draft.get("staff_count", 0))
staff_count = st.number_input("თანამშრომლების რაოდენობა", min_value=0, max_value=30, value=saved_staff_count)

draft_staff = draft.get("staff", [])
staff_members = []
for si in range(int(staff_count)):
    st.markdown(f"**თანამშრომელი #{si+1}**")
    sc1, sc2, sc3, sc4, sc5 = st.columns([2, 2, 2, 2, 3])
    st_prev = draft_staff[si] if si < len(draft_staff) else {}

    s_comp = sc1.text_input("კომპანია", value=st_prev.get("comp", ""), key=f"st_c_{si}")
    s_pos = sc2.text_input("პოზიცია", value=st_prev.get("pos", ""), key=f"st_p_{si}")
    s_name = sc3.text_input("სახელი და გვარი", value=st_prev.get("name", ""), key=f"st_n_{si}")
    s_phone = sc4.text_input("ტელ. ნომერი", value=st_prev.get("phone", ""), key=f"st_ph_{si}")
    s_mail = sc5.text_input("ელ. ფოსტა", value=st_prev.get("mail", ""), key=f"st_m_{si}")
    staff_members.append({"num": si+1, "comp": s_comp, "pos": s_pos, "name": s_name, "phone": s_phone, "mail": s_mail})

current_state_to_save = {
    "rep_num": rep_num,
    "rep_date": rep_date.strftime("%Y-%m-%d"),
    "prepared_by": prepared_by,
    "custom_doc_code": custom_doc_code,
    "client_name": client_name,
    "site_location": site_location,
    "t9": t9, "w9": w9, "cond9": cond9,
    "t14": t14, "w14": w14, "cond14": cond14,
    "t18": t18, "w18": w18, "cond18": cond18,
    "task_count": int(task_count),
    "tasks": tasks,
    "photo_slots_count": int(photo_slots_count),
    "captions": draft_caps,
    "staff_count": int(staff_count),
    "staff": staff_members
}
save_draft(current_state_to_save)

st.write("")
generate_btn = st.button("🚀 რეპორტის შექმნა / Generate DPR", type="primary", use_container_width=True)

# ==================== DOCUMENT GENERATION ====================
if generate_btn:
    try:
        doc = docx.Document(target_file)
        weather_icons = load_weather_icons(target_file)
        formatted_date = rep_date.strftime("%Y/%m/%d")
        final_doc_code = custom_doc_code.strip() if custom_doc_code.strip() else f"CMC-CMS-DPR-{rep_num or '001'}-{rep_date.strftime('%Y%m%d')}-Rev00"

        replacements = {
            "{{ doc_code }}": final_doc_code,
            "{{ rep_num }}": rep_num if rep_num else "-",
            "{{ rep_date }}": formatted_date,
            "{{ client }}": client_name if client_name else "-",
            "{{ location }}": site_location if site_location else "-",
            "{{ prepared_by }}": prepared_by if prepared_by else "-",
            "{{ t9 }}": t9 if t9 else "-",
            "{{ cond9 }}": cond9 if cond9 else "",
            "{{ w9 }}": w9 if w9 else "-",
            "{{ t14 }}": t14 if t14 else "-",
            "{{ cond14 }}": cond14 if cond14 else "",
            "{{ w14 }}": w14 if w14 else "-",
            "{{ t18 }}": t18 if t18 else "-",
            "{{ cond18 }}": cond18 if cond18 else "",
            "{{ w18 }}": w18 if w18 else "-",
            "007": rep_num if rep_num else "-",
            "2026/10/08": formatted_date,
            "მურად ახალაძე": prepared_by if prepared_by else "-",
            "სამუშაეობის": "სამუშაოების",
            "სამუშაეობა": "სამუშაოები"
        }

        # 1. Update Paragraphs
        for p in doc.paragraphs:
            for k, v in replacements.items():
                safe_replace(p, k, v)

        # 2. Iterate Tables
        for table in doc.tables:
            # A. General Placeholder Replacements
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for k, v in replacements.items():
                            safe_replace(p, k, v)

            # B. Weather Row: Clean Wind Speed & Weather Icons (Eliminate broken IF/REF field codes)
            for row in table.rows:
                row_txt = " ".join([c.text for c in row.cells]).lower()
                if "wind speed" in row_txt or "ქარის სიჩქარე" in row_txt:
                    if len(row.cells) == 6:
                        set_clean_wind_cell(row.cells[0], w9)
                        set_weather_icon_cell(row.cells[1], cond9, weather_icons)
                        set_clean_wind_cell(row.cells[2], w14)
                        set_weather_icon_cell(row.cells[3], cond14, weather_icons)
                        set_clean_wind_cell(row.cells[4], w18)
                        set_weather_icon_cell(row.cells[5], cond18, weather_icons)
                    else:
                        seen_tcs = set()
                        wind_cells = []
                        for c in row.cells:
                            if c._tc not in seen_tcs and ("wind speed" in c.text.lower() or "ქარის სიჩქარე" in c.text):
                                seen_tcs.add(c._tc)
                                wind_cells.append(c)
                        if len(wind_cells) >= 1:
                            set_clean_wind_cell(wind_cells[0], w9)
                        if len(wind_cells) >= 2:
                            set_clean_wind_cell(wind_cells[1], w14)
                        if len(wind_cells) >= 3:
                            set_clean_wind_cell(wind_cells[2], w18)

            # C. DAILY WORKS TABLE: 100% Centered Population
            works_h_idx = None
            photo_b_idx = None
            for r_idx, row in enumerate(table.rows):
                r_txt = " ".join([c.text for c in row.cells])
                if ("WBS" in r_txt or "Code" in r_txt) and ("შემსრულებელი" in r_txt or "Contractor" in r_txt or "სამუშაოს დასახელება" in r_txt):
                    works_h_idx = r_idx
                if works_h_idx is not None and r_idx > works_h_idx and ("ფოტომასალა" in r_txt or "Work Progress Photos" in r_txt or "სამუშაო პროცესი" in r_txt):
                    photo_b_idx = r_idx
                    break

            if works_h_idx is not None:
                for c in table.rows[works_h_idx].cells:
                    try:
                        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    except Exception:
                        pass
                    for p in c.paragraphs:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

                sample_task_tr = None
                task_rows_to_delete = []

                if photo_b_idx is not None:
                    photo_banner_row = table.rows[photo_b_idx]
                    for r_i in range(works_h_idx + 1, photo_b_idx):
                        if sample_task_tr is None:
                            sample_task_tr = copy.deepcopy(table.rows[r_i]._tr)
                        task_rows_to_delete.append(table.rows[r_i])
                    if sample_task_tr is None:
                        sample_task_tr = copy.deepcopy(table.rows[works_h_idx]._tr)

                    for r in task_rows_to_delete:
                        table._tbl.remove(r._tr)

                    for tsk in tasks:
                        new_tr = copy.deepcopy(sample_task_tr)
                        new_row = docx.table._Row(new_tr, table)
                        vals = [
                            str(tsk["num"]),
                            str(tsk["contractor"]),
                            str(tsk["manpower"]),
                            str(tsk["wbs"]),
                            str(tsk["name"]),
                            str(tsk["desc"])
                        ]
                        for c_idx, val in enumerate(vals):
                            if c_idx < len(new_row.cells):
                                populate_cell(new_row.cells[c_idx], val, font_size=8.0, bold=(c_idx in [0, 4]), align=WD_ALIGN_PARAGRAPH.CENTER)
                        photo_banner_row._tr.addprevious(new_tr)
                else:
                    if len(table.rows) > works_h_idx + 1:
                        sample_task_tr = copy.deepcopy(table.rows[works_h_idx + 1]._tr)
                    else:
                        sample_task_tr = copy.deepcopy(table.rows[works_h_idx]._tr)

                    while len(table.rows) > works_h_idx + 1:
                        table._tbl.remove(table.rows[-1]._tr)

                    for tsk in tasks:
                        new_tr = copy.deepcopy(sample_task_tr)
                        new_row = docx.table._Row(new_tr, table)
                        vals = [
                            str(tsk["num"]),
                            str(tsk["contractor"]),
                            str(tsk["manpower"]),
                            str(tsk["wbs"]),
                            str(tsk["name"]),
                            str(tsk["desc"])
                        ]
                        for c_idx, val in enumerate(vals):
                            if c_idx < len(new_row.cells):
                                populate_cell(new_row.cells[c_idx], val, font_size=8.0, bold=(c_idx in [0, 4]), align=WD_ALIGN_PARAGRAPH.CENTER)
                        table._tbl.append(new_tr)

            # D. PHOTOS TABLE: Centered Pictures
            t_text = " ".join([c.text for row in table.rows for c in row.cells])
            if "ფოტომასალა" in t_text or "Daily Work Progress Photos" in t_text:
                num_uploaded = len(photos_data)
                wp_idx = None
                for r_idx, row in enumerate(table.rows):
                    r_txt = " ".join([c.text for c in row.cells])
                    if "სამუშაო პროცესი" in r_txt or "Work Process" in r_txt:
                        wp_idx = r_idx
                        break

                if wp_idx is not None:
                    pairs = [
                        (wp_idx + 1, wp_idx + 2),
                        (wp_idx + 3, wp_idx + 4),
                        (wp_idx + 5, wp_idx + 6)
                    ]

                    for p_i, (p_row_idx, c_row_idx) in enumerate(pairs):
                        if p_row_idx < len(table.rows) and c_row_idx < len(table.rows):
                            p_row = table.rows[p_row_idx]
                            c_row = table.rows[c_row_idx]

                            slot_a = p_i * 2
                            slot_b = p_i * 2 + 1

                            if slot_a < num_uploaded:
                                insert_centered_picture(p_row.cells[0], photos_data[slot_a]["file"].getvalue(), width=Inches(3.15))
                                populate_cell(c_row.cells[0], photos_data[slot_a]["caption"], font_size=8.0, align=WD_ALIGN_PARAGRAPH.CENTER)
                            else:
                                clear_cell_completely(p_row.cells[0])
                                clear_cell_completely(c_row.cells[0])

                            if slot_b < num_uploaded:
                                insert_centered_picture(p_row.cells[1], photos_data[slot_b]["file"].getvalue(), width=Inches(3.15))
                                populate_cell(c_row.cells[1], photos_data[slot_b]["caption"], font_size=8.0, align=WD_ALIGN_PARAGRAPH.CENTER)
                            else:
                                clear_cell_completely(p_row.cells[1])
                                clear_cell_completely(c_row.cells[1])

                    # Extra Photos (> 6)
                    if num_uploaded > 6:
                        remaining_photos = photos_data[6:]
                        proc_counter = 4
                        sample_header_row = table.rows[wp_idx]
                        sample_photo_row = table.rows[wp_idx + 1]
                        sample_cap_row = table.rows[wp_idx + 2]

                        hs_row = None
                        for r_idx in range(wp_idx + 7, len(table.rows)):
                            r_txt = " ".join([c.text for c in table.rows[r_idx].cells])
                            if "Health & Safety" in r_txt or "უსაფრთხოების" in r_txt:
                                hs_row = table.rows[r_idx]
                                break

                        for pair_start in range(0, len(remaining_photos), 2):
                            new_h_tr = copy.deepcopy(sample_header_row._tr)
                            new_h_row = docx.table._Row(new_h_tr, table)
                            t_cell = new_h_row.cells[0]
                            if len(new_h_row.cells) > 1:
                                t_cell.merge(new_h_row.cells[1])
                            t_cell.text = f"სამუშაო პროცესი / Work Process №{proc_counter}"
                            for r in t_cell.paragraphs[0].runs:
                                r.font.bold = True
                            if proc_counter == 4 or (proc_counter - 4) % 3 == 0:
                                t_cell.paragraphs[0].paragraph_format.page_break_before = True

                            new_p_tr = copy.deepcopy(sample_photo_row._tr)
                            new_p_row = docx.table._Row(new_p_tr, table)
                            new_c_tr = copy.deepcopy(sample_cap_row._tr)
                            new_c_row = docx.table._Row(new_c_tr, table)

                            insert_centered_picture(new_p_row.cells[0], remaining_photos[pair_start]["file"].getvalue(), width=Inches(3.15))
                            populate_cell(new_c_row.cells[0], remaining_photos[pair_start]["caption"], font_size=8.0, align=WD_ALIGN_PARAGRAPH.CENTER)

                            if pair_start + 1 < len(remaining_photos):
                                insert_centered_picture(new_p_row.cells[1], remaining_photos[pair_start + 1]["file"].getvalue(), width=Inches(3.15))
                                populate_cell(new_c_row.cells[1], remaining_photos[pair_start + 1]["caption"], font_size=8.0, align=WD_ALIGN_PARAGRAPH.CENTER)
                            else:
                                clear_cell_completely(new_p_row.cells[1])
                                clear_cell_completely(new_c_row.cells[1])

                            if hs_row is not None:
                                hs_row._tr.addprevious(new_h_tr)
                                hs_row._tr.addprevious(new_p_tr)
                                hs_row._tr.addprevious(new_c_tr)
                            else:
                                table._tbl.append(new_h_tr)
                                table._tbl.append(new_p_tr)
                                table._tbl.append(new_c_tr)

                            proc_counter += 1

            # E. KEY STAFF TABLE: 100% Centered
            staff_h_idx = None
            sig_r_idx = None
            for r_idx, row in enumerate(table.rows):
                r_txt = " ".join([c.text for c in row.cells])
                if ("კომპანია" in r_txt or "Company" in r_txt) and ("E-mail" in r_txt or "ელ. ფოსტა" in r_txt or "Phone number" in r_txt):
                    staff_h_idx = r_idx
                if staff_h_idx is not None and r_idx > staff_h_idx and ("ვადასტურებ" in r_txt or "განვიხილე" in r_txt or "შენიშვნა" in r_txt):
                    sig_r_idx = r_idx
                    break

            if staff_h_idx is not None:
                for c in table.rows[staff_h_idx].cells:
                    try:
                        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    except Exception:
                        pass
                    for p in c.paragraphs:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

                sample_staff_tr = None
                staff_rows_to_delete = []

                if sig_r_idx is not None:
                    sig_row = table.rows[sig_r_idx]
                    for r_i in range(staff_h_idx + 1, sig_r_idx):
                        if sample_staff_tr is None:
                            sample_staff_tr = copy.deepcopy(table.rows[r_i]._tr)
                        staff_rows_to_delete.append(table.rows[r_i])
                    if sample_staff_tr is None:
                        sample_staff_tr = copy.deepcopy(table.rows[staff_h_idx]._tr)

                    for r in staff_rows_to_delete:
                        table._tbl.remove(r._tr)

                    if staff_members:
                        for sm in staff_members:
                            new_tr = copy.deepcopy(sample_staff_tr)
                            new_row = docx.table._Row(new_tr, table)
                            vals = [
                                str(sm["num"]),
                                str(sm["comp"]),
                                str(sm["pos"]),
                                str(sm["name"]),
                                str(sm["phone"]),
                                str(sm["mail"])
                            ]
                            for c_idx, val in enumerate(vals):
                                if c_idx < len(new_row.cells):
                                    populate_cell(new_row.cells[c_idx], val, font_size=7.5, bold=(c_idx in [0, 3]), align=WD_ALIGN_PARAGRAPH.CENTER)
                            sig_row._tr.addprevious(new_tr)
                    else:
                        new_tr = copy.deepcopy(sample_staff_tr)
                        new_row = docx.table._Row(new_tr, table)
                        for c in new_row.cells:
                            clear_cell_completely(c)
                        sig_row._tr.addprevious(new_tr)
                else:
                    if len(table.rows) > staff_h_idx + 1:
                        sample_staff_tr = copy.deepcopy(table.rows[staff_h_idx + 1]._tr)
                    else:
                        sample_staff_tr = copy.deepcopy(table.rows[staff_h_idx]._tr)

                    while len(table.rows) > staff_h_idx + 1:
                        table._tbl.remove(table.rows[-1]._tr)

                    if staff_members:
                        for sm in staff_members:
                            new_tr = copy.deepcopy(sample_staff_tr)
                            new_row = docx.table._Row(new_tr, table)
                            vals = [
                                str(sm["num"]),
                                str(sm["comp"]),
                                str(sm["pos"]),
                                str(sm["name"]),
                                str(sm["phone"]),
                                str(sm["mail"])
                            ]
                            for c_idx, val in enumerate(vals):
                                if c_idx < len(new_row.cells):
                                    populate_cell(new_row.cells[c_idx], val, font_size=7.5, bold=(c_idx in [0, 3]), align=WD_ALIGN_PARAGRAPH.CENTER)
                            table._tbl.append(new_tr)
                    else:
                        new_tr = copy.deepcopy(sample_staff_tr)
                        new_row = docx.table._Row(new_tr, table)
                        for c in new_row.cells:
                            clear_cell_completely(c)
                        table._tbl.append(new_tr)

        bio = io.BytesIO()
        doc.save(bio)
        st.session_state["docx_bytes"] = bio.getvalue()
        st.session_state["doc_code"] = final_doc_code

        # ==================== PDF CONVERSION (Franky1 Streamlit-docx-converter) ====================
        temp_dir = Path(tempfile.mkdtemp())
        doc_path = temp_dir / f"{final_doc_code}.docx"
        with open(doc_path, "wb") as f_tmp:
            f_tmp.write(bio.getvalue())

        output_pdf, conv_exc = convert_doc_to_pdf_native(doc_path, output_dir=temp_dir, timeout=60)

        # Fallback for Windows if soffice is not in PATH
        if output_pdf is None and os.name == 'nt':
            try:
                import pythoncom
                pythoncom.CoInitialize()
                import win32com.client
                word_app = win32com.client.DispatchEx("Word.Application")
                word_app.Visible = False
                word_app.DisplayAlerts = False
                pdf_target = temp_dir / f"{final_doc_code}.pdf"
                try:
                    w_doc = word_app.Documents.Open(str(doc_path.resolve()))
                    w_doc.ExportAsFixedFormat(str(pdf_target.resolve()), 17)
                    w_doc.Close(False)
                    if pdf_target.exists():
                        output_pdf = pdf_target
                finally:
                    word_app.Quit()
                    pythoncom.CoUninitialize()
            except Exception as e_word:
                if conv_exc is None:
                    conv_exc = e_word

        if output_pdf and output_pdf.exists() and output_pdf.stat().st_size > 0:
            with open(output_pdf, "rb") as f_pdf:
                st.session_state["pdf_bytes"] = f_pdf.read()
            pdf_converted = True
        else:
            pdf_converted = False
            conversion_error = conv_exc

        try:
            shutil.rmtree(temp_dir, ignore_errors=True)
        except Exception:
            pass

        if not pdf_converted:
            st.warning(f"💡 შენიშვნა: PDF ვერ დაგენერირდა ({conversion_error}). გადმოწერეთ Word ფაილი.")

    except Exception as e:
        st.error(f"შეცდომა რეპორტის შექმნისას: {e}")

# ==================== DOWNLOAD BUTTONS ====================
if st.session_state["docx_bytes"] is not None:
    d_code = st.session_state["doc_code"]
    st.success(f"🎉 რეპორტი მზადაა: **{d_code}**")

    b_col1, b_col2 = st.columns(2)
    b_col1.download_button(
        label="📥 1. გადმოწერეთ Word (.docx)",
        data=st.session_state["docx_bytes"],
        file_name=f"{d_code}.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        use_container_width=True
    )

    if st.session_state["pdf_bytes"] is not None:
        b_col2.download_button(
            label="🖨️ 2. გადმოწერეთ დასაბეჭდი PDF (.pdf)",
            data=st.session_state["pdf_bytes"],
            file_name=f"{d_code}.pdf",
            mime="application/pdf",
            use_container_width=True
        )
    else:
        b_col2.info("💡 PDF ღილაკისთვის Streamlit Cloud-ზე GitHub რეპოზიტორიაში დაამატეთ `packages.txt` ჩანაწერით: `libreoffice`.")
